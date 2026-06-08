#!/usr/bin/env python
"""
Build a rankings table from main_slim_normalized results.

Section 1 — last-generation metrics (train RMSE, test RMSE, size, M_phi):
  rank per algorithm per dataset + median rank across datasets.

Section 2 — post-simplification metrics (ell_after, m_phi_after):
  same structure, drawn from results_normalized_simplification.csv.

Saves:
  log/figs/rankings_table.csv   — flat CSV with multi-level header
  log/figs/rankings_table.xlsx  — Excel with conditional formatting (optional)
"""
import os
import sys

import numpy as np
import pandas as pd

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

LOG_DIR = os.path.join(os.path.dirname(__file__), 'log')
GEN_LOG = os.path.join(LOG_DIR, 'results_normalized_generations.csv')
OUT_DIR = os.path.join(LOG_DIR, 'figs')
os.makedirs(OUT_DIR, exist_ok=True)

# ── Config ────────────────────────────────────────────────────────────────────
ALGOS = [
    'SLIM+2SIG', 'SLIM*2SIG',
    'SLIM+ABS',  'SLIM*ABS',
    'SLIM+1SIG', 'SLIM*1SIG',
    'SLIM+NORM1','SLIM*NORM1',
    'SLIM+NORM2','SLIM*NORM2',
]
DATASETS = ['concrete', 'energy', 'instanbul', 'ppb', 'resid_build_sale_price', 'toxicity']
DS_SHORT  = {
    'concrete':              'Concrete',
    'energy':                'Energy',
    'instanbul':             'Istanbul',
    'ppb':                   'PPB',
    'resid_build_sale_price':'Resid.',
    'toxicity':              'Toxicity',
}

# (column, display label, ascending)  ascending=True → lower value = rank 1
METRICS = [
    ('train_fitness', 'Train RMSE', True),
    ('test_fitness',  'Test RMSE',  True),
    ('nodes_count',   'Size',       True),
    ('m_phi',         'M_phi',      False),  # higher = better
]

# ── Load last generation ──────────────────────────────────────────────────────
print('Loading data ...', flush=True)
usecols = ['algo', 'dataset', 'seed', 'generation',
           'train_fitness', 'test_fitness', 'nodes_count', 'm_phi']
df = pd.read_csv(GEN_LOG, usecols=usecols)
last_gen = int(df['generation'].max())
df_last  = df[df['generation'] == last_gen]
print(f'  Using generation {last_gen}  ({len(df_last)} rows)', flush=True)

# ── Per-(algo, dataset) medians ───────────────────────────────────────────────
metric_cols = [m[0] for m in METRICS]
medians = (df_last
           .groupby(['algo', 'dataset'])[metric_cols]
           .median())   # MultiIndex: (algo, dataset)

# ── Build table ───────────────────────────────────────────────────────────────
# Columns will be a MultiIndex: (metric_label, dataset_short | 'Median')
col_tuples = []
data_dict  = {}

for col, label, ascending in METRICS:
    # ------ per-dataset ranks -----------------------------------------------
    for ds in DATASETS:
        ds_label = DS_SHORT[ds]
        col_key  = (label, ds_label)
        col_tuples.append(col_key)
        # slice this dataset for all algos, rank
        ds_med   = medians.xs(ds, level='dataset')[col]   # Series indexed by algo
        ranks    = ds_med.rank(ascending=ascending, method='min').astype(int)
        data_dict[col_key] = {a: int(ranks[a]) if a in ranks.index else np.nan
                              for a in ALGOS}

    # ------ median rank across datasets ------------------------------------
    col_key = (label, 'Median rank')
    col_tuples.append(col_key)
    data_dict[col_key] = {}
    for algo in ALGOS:
        per_ds_ranks = [data_dict[(label, DS_SHORT[ds])][algo]
                        for ds in DATASETS
                        if (label, DS_SHORT[ds]) in data_dict
                        and algo in data_dict[(label, DS_SHORT[ds])]]
        data_dict[col_key][algo] = float(np.median(per_ds_ranks)) if per_ds_ranks else np.nan

# ── Post-simplification metrics ───────────────────────────────────────────────
SIMP_LOG = os.path.join(LOG_DIR, 'results_normalized_simplification.csv')
SIMP_METRICS = [
    ('ell_after',   'Size (simp)',  True),
    ('m_phi_after', 'M_phi (simp)', False),
]

simp_df = pd.read_csv(SIMP_LOG,
                      usecols=['algo', 'dataset', 'seed',
                                'ell_before', 'm_phi_before',
                                'ell_after', 'm_phi_after'])
# Analysis-side filter: if SymPy made things worse, use the original
simp_df['ell_after']   = simp_df[['ell_after',   'ell_before'  ]].min(axis=1)
simp_df['m_phi_after'] = simp_df[['m_phi_after', 'm_phi_before']].max(axis=1)

simp_medians = (simp_df
                .groupby(['algo', 'dataset'])[['ell_after', 'm_phi_after']]
                .median())

for col, label, ascending in SIMP_METRICS:
    for ds in DATASETS:
        ds_label = DS_SHORT[ds]
        col_key  = (label, ds_label)
        col_tuples.append(col_key)
        ds_med = simp_medians.xs(ds, level='dataset')[col]
        ranks  = ds_med.rank(ascending=ascending, method='min').astype(int)
        data_dict[col_key] = {a: int(ranks[a]) if a in ranks.index else np.nan
                              for a in ALGOS}

    col_key = (label, 'Median rank')
    col_tuples.append(col_key)
    data_dict[col_key] = {}
    for algo in ALGOS:
        per_ds_ranks = [data_dict[(label, DS_SHORT[ds])][algo]
                        for ds in DATASETS
                        if (label, DS_SHORT[ds]) in data_dict
                        and algo in data_dict[(label, DS_SHORT[ds])]]
        data_dict[col_key][algo] = float(np.median(per_ds_ranks)) if per_ds_ranks else np.nan

# Rebuild table with simp columns added
table = pd.DataFrame(
    {k: pd.Series(v) for k, v in data_dict.items()},
    index=ALGOS
)
table.columns = pd.MultiIndex.from_tuples(col_tuples)
table.index.name = 'Algorithm'

ALL_METRICS = METRICS + SIMP_METRICS

# ── Print ─────────────────────────────────────────────────────────────────────
pd.set_option('display.max_columns', None)
pd.set_option('display.width', 220)
pd.set_option('display.float_format', '{:.3f}'.format)

print('\n' + '='*80)
for col, label, ascending in ALL_METRICS:
    sub = table[label]
    rank_cols = [DS_SHORT[ds] for ds in DATASETS]
    print(f'\n--- {label} ({"lower" if ascending else "higher"} = better, rank 1 = best) ---')
    print(sub[rank_cols + ['Median rank']].to_string())
print('='*80)

# ── Save CSV ──────────────────────────────────────────────────────────────────
csv_path = os.path.join(OUT_DIR, 'rankings_table.csv')
table.to_csv(csv_path)
print(f'\nSaved -> {csv_path}')

# ── Save Excel with rank-based colour gradient ────────────────────────────────
try:
    import openpyxl  # noqa: F401 — just checking it's available

    xlsx_path = os.path.join(OUT_DIR, 'rankings_table.xlsx')

    # Flatten to a single-row header for Excel
    flat = table.copy()
    flat.columns = [f'{m} | {d}' for m, d in flat.columns]
    flat = flat.reset_index()

    with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
        flat.to_excel(writer, index=False, sheet_name='Rankings')
        ws = writer.sheets['Rankings']

        # Auto-fit column widths
        for col_cells in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = max_len + 2

        # Colour rank columns: green=1, red=10
        from openpyxl.styles import PatternFill
        import colorsys

        def rank_colour(rank, n=10):
            if pd.isna(rank):
                return 'FFFFFF'
            t = (int(rank) - 1) / (n - 1)      # 0=best, 1=worst
            r = int(255 * t + 40 * (1 - t))     # red channel
            g = int(200 * (1 - t) + 40 * t)     # green channel
            b = 60
            return f'{r:02X}{g:02X}{b:02X}'

        header_row = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                col_header = header_row[cell.column - 1]
                if col_header and '| Median rank' not in str(col_header) and col_header != 'Algorithm':
                    try:
                        hex_c = rank_colour(int(cell.value))
                        cell.fill = PatternFill(fill_type='solid', fgColor=hex_c)
                    except (TypeError, ValueError):
                        pass

    print(f'Saved -> {xlsx_path}')
except ImportError:
    print('openpyxl not installed — skipping Excel output (pip install openpyxl)')

print('\nDone.')
